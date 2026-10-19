#!/usr/bin/env python3
"""Fill the WISE Digital Audit template from an audit JSON file.

Usage:
    python3 fill_audit.py audit.json OUTPUT.xlsx [--template assets/audit-template.csv]

Writes OUTPUT.xlsx with two sheets: "WISE Digital Audit" in the template's exact
layout (also written as OUTPUT.csv next to it) and "How We Got These Results",
one row per cell with the method used, the evidence recorded in audit.json
(`evidence`, `sources`, directory `detail`, keyword `reason`, ranking `note`,
`observation_basis`) and the source. Cells are located by their labels in the
template, not by fixed coordinates, so small template edits do not break the
filler. Exits non-zero with a list of problems if the JSON is incomplete or
inconsistent.
"""
import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

DIRECTORY_TOTAL = 21

# Left and right result tables: label in the template -> JSON key.
# The result sits two columns to the right of the label.
SEO_ROWS = {
    "Website Platform": "website_platform",
    "URL Structure": "url_structure",
    "Meta Titles": "meta_titles",
    "Meta Descriptions": "meta_descriptions",
    "Headers / H Tags": "headers",
    "Structured Data": "structured_data",
    "Page Word Count": "page_word_count",
    "New Content Monthly": "new_content_monthly",
    "Duplicate Websites": "duplicate_websites",
    "Duplicate Content": "duplicate_content",
    "Backlinks": "backlinks",
    "Internal Link Building": "internal_link_building",
    "External Link Building": "external_link_building",
    "Domain Authority": "domain_authority",
    "Domain Age": "domain_age",
}
UI_ROWS = {
    "SSL Secured": "ssl_secured",
    "ADA Compliant / ADA Software": "ada_compliant",
    "Privacy/Consent Compliance": "privacy_consent",
    "Tone & Voice": "tone_and_voice",
    "Original Content (Photos, Videos, etc)": "original_content",
    "Clear Customer Journey": "clear_customer_journey",
    "Clear Conversion Path (CTAs)": "clear_conversion_path",
    "Overall Test Results": "cwv_overall",
    "- Desktop Page Optimization": "desktop_optimization",
    "- Desktop Speed (Speed Index in seconds)": "desktop_speed_seconds",
    "- Desktop LCP (seconds)": "desktop_lcp_seconds",
    "- Mobile Page Optimization": "mobile_optimization",
    "- Mobile Speed (4G Speed Index in seconds)": "mobile_speed_seconds",
    "- Mobile LCP (4G seconds)": "mobile_lcp_seconds",
}
# Header label -> (section, key). The value goes in the cell directly below.
UNDER_HEADER = {
    "# of Core Pages": ("content", "core_pages"),
    "# of Blog Posts": ("content", "blog_posts"),
    "AI-Written Blogs": ("content", "ai_written_blogs"),
    "Google (Gemini)": ("ai_visibility", "gemini"),
    "ChatGPT": ("ai_visibility", "chatgpt"),
    "Google Analytics": ("tracking", "google_analytics"),
    "Google Tag Manager": ("tracking", "google_tag_manager"),
    "# of Business Locations": ("tracking", "business_locations"),
    "Google Ads": ("tracking", "google_ads"),
    "Meta Ads": ("tracking", "meta_ads"),
    "Other Ads": ("tracking", "other_ads"),
}
REVIEW_COLUMNS = {"Google Business Profile": "gbp", "Yelp": "yelp", "Facebook": "facebook"}

RED = PatternFill("solid", fgColor="F8D7DA")
GREEN = PatternFill("solid", fgColor="D4EDDA")
BAD = {"incorrect", "no", "fail"}
GOOD = {"correct", "yes", "pass"}
# For these checks "Yes" is the bad answer.
INVERTED = {"Duplicate Websites", "Duplicate Content", "AI-Written Blogs"}


def norm(text):
    return " ".join(str(text).split()).rstrip(":").strip()


def load_grid(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    width = max(len(r) for r in rows)
    return [r + [""] * (width - len(r)) for r in rows]


def find(grid, label, start_row=0, col=None):
    want = norm(label).lower()
    for r in range(start_row, len(grid)):
        for c, cell in enumerate(grid[r]):
            if (col is None or c == col) and norm(cell).lower() == want:
                return r, c
    return None


def put(grid, r, c, value):
    while c >= len(grid[r]):
        grid[r].append("")
    grid[r][c] = "" if value is None else str(value)


def fill(grid, data, problems):
    def need(section, key):
        sec = data.get(section, {})
        if key not in sec or sec[key] in (None, ""):
            problems.append(f"missing {section}.{key}")
            return ""
        return sec[key]

    # Header: date and company.
    hit = next(((r, c) for r in range(len(grid)) for c, v in enumerate(grid[r])
                if v.strip().upper().startswith("DIGITAL AUDIT")), None)
    if hit:
        put(grid, *hit, f"DIGITAL AUDIT  {need('prospect', 'audit_date')}")
    hit = find(grid, "[Company Name]")
    if hit:
        put(grid, *hit, need("prospect", "company_name"))
    hit = find(grid, "[url]")
    if hit:
        put(grid, *hit, need("prospect", "url"))

    for table, section in ((SEO_ROWS, "seo"), (UI_ROWS, "ui")):
        for label, key in table.items():
            hit = find(grid, label)
            if not hit:
                problems.append(f"template label not found: {label}")
                continue
            put(grid, hit[0], hit[1] + 2, need(section, key))

    for label, (section, key) in UNDER_HEADER.items():
        hit = find(grid, label)
        if not hit:
            problems.append(f"template label not found: {label}")
            continue
        put(grid, hit[0] + 1, hit[1], need(section, key))
    # "Other" appears twice (AI visibility and ads); the AI one is in the Gemini row.
    gem = find(grid, "Google (Gemini)")
    if gem:
        for c in range(gem[1], len(grid[gem[0]])):
            if norm(grid[gem[0]][c]) == "Other":
                put(grid, gem[0] + 1, c, need("ai_visibility", "other"))

    # Reviews: stars row then review-count row.
    for label, key in REVIEW_COLUMNS.items():
        hit = find(grid, label)
        if not hit:
            problems.append(f"template label not found: {label}")
            continue
        rev = data.get("reviews", {}).get(key, {})
        stars, count = rev.get("stars"), rev.get("count")
        if stars is None or count is None:
            problems.append(f"missing reviews.{key}.stars/count")
        if str(stars).lower() == "unknown" or str(count).lower() == "unknown":
            put(grid, hit[0] + 1, hit[1], "Unknown")
            put(grid, hit[0] + 2, hit[1], "Unknown")
            continue
        put(grid, hit[0] + 1, hit[1], f"{stars if stars is not None else 0} - Stars")
        put(grid, hit[0] + 2, hit[1], f"{count if count is not None else 0} - Reviews")

    # Directory listings.
    d = data.get("directories", {})
    inaccurate, not_found = d.get("inaccurate"), d.get("not_found")
    if inaccurate is None or not_found is None:
        problems.append("missing directories.inaccurate/not_found")
        inaccurate, not_found = inaccurate or 0, not_found or 0
    if inaccurate + not_found > DIRECTORY_TOTAL:
        problems.append(f"directories: inaccurate + not_found exceeds {DIRECTORY_TOTAL}")
    rate = round(100 * (inaccurate + not_found) / DIRECTORY_TOTAL)
    for label, value in (("Inaccurate", f"{inaccurate}/{DIRECTORY_TOTAL}"),
                         ("Not Found", f"{not_found}/{DIRECTORY_TOTAL}"),
                         ("Error Rate", f"{rate}%")):
        hit = find(grid, label)
        if hit:
            put(grid, hit[0] + 1, hit[1], value)

    # Keyword groups (5 rows under the header) and total.
    groups = data.get("keyword_groups", [])
    if len(groups) != 5:
        problems.append(f"keyword_groups must have 5 entries, got {len(groups)}")
    hdr = find(grid, "Keyword Groups")
    total = 0
    if hdr:
        for i, g in enumerate(groups[:5]):
            r = hdr[0] + 1 + i
            put(grid, r, hdr[1], g.get("keyword", ""))
            put(grid, r, hdr[1] + 1, g.get("searches_per_month", 0))
            put(grid, r, hdr[1] + 2, g.get("city", ""))
            total += int(g.get("searches_per_month") or 0)
    hit = find(grid, "Total Monthly Searches")
    if hit:
        put(grid, hit[0], hit[1] + 1, total)

    # Observations: prospect-facing only, one per row under the header.
    obs = data.get("observations", [])
    if not 3 <= len(obs) <= 5:
        problems.append(f"observations must have 3-5 entries, got {len(obs)}")
    hdr = find(grid, "Key Observations & Concerns")
    if hdr:
        for i, text in enumerate(obs[:5]):
            put(grid, hdr[0] + 1 + i, hdr[1], text)

    # Rankings breakdown: 25 rows under the "Keywords" header.
    ranks = data.get("rankings", [])
    if len(ranks) != 25:
        problems.append(f"rankings must have 25 entries, got {len(ranks)}")
    hdr = find(grid, "Keywords", start_row=find(grid, "Search Rankings Breakdown")[0]
               if find(grid, "Search Rankings Breakdown") else 0)
    maps_hits = organic_hits = 0
    if hdr:
        maps_col = find(grid, "Top 3 (Maps)", start_row=hdr[0])[1]
        org_col = find(grid, "1st Page (Organic)", start_row=hdr[0])[1]
        for i, row in enumerate(ranks[:25]):
            r = hdr[0] + 1 + i
            put(grid, r, hdr[1], row.get("keyword", ""))
            put(grid, r, hdr[1] + 1, row.get("city", ""))
            m, o = row.get("maps_position"), row.get("organic_position")
            put(grid, r, maps_col, m if m else "-")
            put(grid, r, org_col, o if o else "-")
            maps_hits += 1 if m and int(m) <= 3 else 0
            organic_hits += 1 if o and int(o) <= 10 else 0
    # The template puts each count in the first filled cell below its label.
    for label, value in (("Top 3 Maps Rankings", maps_hits), ("Top 10 Organic Rankings", organic_hits)):
        hit = find(grid, label)
        if hit:
            r, c = hit
            target = next((rr for rr in range(r + 1, min(r + 4, len(grid)))
                           if c < len(grid[rr]) and grid[rr][c].strip()), r + 1)
            put(grid, target, c, value)


def colour(label, value):
    v = norm(value).lower()
    if not v:
        return None
    bad, good = (GOOD, BAD) if label in INVERTED else (BAD, GOOD)
    if v in bad:
        return RED
    if v in good:
        return GREEN
    return None


# Second sheet: how each result was reached. One fixed method note per cell, plus the
# page-level evidence and source the model recorded in audit.json under `evidence`
# (string, or {"note": ..., "source": ...}) and `sources`.
METHODS = {
    "prospect.company_name": "Name on the website, cross-checked with the call transcript.",
    "prospect.url": "URL supplied by the sales manager (or confirmed from the transcript).",
    "prospect.audit_date": "Date the audit was run.",
    "seo.website_platform": "Generator meta tag, asset paths (/wp-content/, cdn.shopify.com, wixstatic, squarespace, webflow) and response headers; transcript if undetectable.",
    "seo.url_structure": "Each key page URL checked for query strings, IDs, uppercase, underscores, file extensions, depth over 3 folders and bare service slugs. Incorrect if any key page fails.",
    "seo.meta_titles": "Site crawl plus page HTML: missing, duplicated across key pages, over 60 characters, or missing the page's main service. Incorrect if any key page fails.",
    "seo.meta_descriptions": "Site crawl plus page HTML: missing, duplicated, under 70 or over 160 characters. Incorrect if any key page fails.",
    "seo.headers": "Page HTML: exactly one H1 naming the main service and no skipped heading levels. Incorrect if any key page fails.",
    "seo.structured_data": "Homepage JSON-LD checked for a LocalBusiness (or subtype) or Organization schema that parses.",
    "seo.page_word_count": "Unique main-content words per key page (header, footer and menu text subtracted). Under 500 fails.",
    "seo.new_content_monthly": "Blog post dates or sitemap lastmod: at least one new post or page in each of the last 3 months.",
    "seo.duplicate_websites": "http, https, www and non-www variants requested and checked for a redirect to one version; business name searched for a second site with the same content.",
    "seo.duplicate_content": "Site crawl duplicate-content issues on key pages.",
    "seo.backlinks": "SE Ranking backlinks summary: 10 or more referring domains is Yes.",
    "seo.internal_link_building": "Page HTML: every core service page linked from at least one other key page besides the menu.",
    "seo.external_link_building": "SE Ranking new and lost referring domains over the last 6 months: growth is Yes.",
    "seo.domain_authority": "SE Ranking Domain Authority (0 to 100 on SE Ranking's own scale).",
    "seo.domain_age": "Registration date from the RDAP record for the domain.",
    "ui.ssl_secured": "HTTPS loads with a valid certificate on the key pages.",
    "ui.ada_compliant": "scripts/check_ada.py on the saved key-page HTML: Pass if an accessibility overlay or widget (accessiBe, UserWay, AudioEye, EqualWeb and similar) is running, or if the quick screen finds no images without alt text, no unlabelled form fields and a lang attribute on every key page. Otherwise Fail. Not a WCAG certification.",
    "ui.privacy_consent": "Privacy policy link in the footer, and a consent banner or consent tool where tracking tags fire.",
    "ui.tone_and_voice": "Five statements scored on the key pages (speaks to the customer, plain language, consistent voice, fits the audience from the call, no spelling or grammar errors). Correct at 3 or more.",
    "ui.original_content": "Key pages checked for own photos or video versus stock-only imagery.",
    "ui.clear_customer_journey": "Every core service reachable from the homepage in one click, and each service page points to a next step.",
    "ui.clear_conversion_path": "Every key page has a visible call or form CTA above the fold on mobile.",
    "ui.cwv_overall": "PageSpeed Insights for the homepage: field-data Core Web Vitals assessment on mobile, or lab thresholds (LCP 2.5 s, CLS 0.1, TBT 200 ms) when there is no field data.",
    "ui.desktop_optimization": "PageSpeed Insights desktop Lighthouse performance score.",
    "ui.desktop_speed_seconds": "PageSpeed Insights desktop lab Speed Index, in seconds.",
    "ui.desktop_lcp_seconds": "PageSpeed Insights desktop Largest Contentful Paint: the 28-day field value (75th percentile) when Google has one, otherwise the lab value, in seconds.",
    "ui.mobile_optimization": "PageSpeed Insights mobile Lighthouse performance score.",
    "ui.mobile_speed_seconds": "PageSpeed Insights mobile lab Speed Index on simulated 4G, in seconds.",
    "ui.mobile_lcp_seconds": "PageSpeed Insights mobile Largest Contentful Paint: the 28-day field value (75th percentile) when Google has one, otherwise the lab value on simulated 4G, in seconds.",
    "content.core_pages": "Sitemap URLs that are service, location, industry, pricing or about pages; blog, legal, offer and thank-you pages excluded.",
    "content.blog_posts": "Post count from the post sitemap.",
    "content.ai_written_blogs": "The 3 most recent posts judged for generic intros, no first-hand detail, listicle structure and stock phrases; Yes if 2 of 3 show 3 or more signs.",
    "ai_visibility.gemini": "SE Ranking AI search data: domain cited or brand mentioned for any prompt in Google AI Overviews or AI Mode.",
    "ai_visibility.chatgpt": "SE Ranking AI search data: domain cited or brand mentioned for any ChatGPT prompt.",
    "ai_visibility.other": "SE Ranking AI search data for Perplexity and any other engine reported.",
    "tracking.google_analytics": "GA4 or Universal Analytics tag found in the page HTML or the Tag Manager container.",
    "tracking.google_tag_manager": "gtm.js container script found in the page HTML.",
    "tracking.business_locations": "Location pages on the site, cross-checked with the transcript.",
    "tracking.google_ads": "AW- conversion tag on the site or in the container, or SE Ranking paid-ads history in the last 3 months. Detection wins over the transcript.",
    "tracking.meta_ads": "Meta Ad Library active ads where reachable; otherwise a Meta pixel plus the prospect saying they run Meta ads.",
    "tracking.other_ads": "Microsoft UET, TikTok, LinkedIn, Nextdoor or similar pixels, or another paid channel named on the call.",
    "reviews.gbp": "Stars and review count read from the Google Business Profile found by searching the business name and city.",
    "reviews.yelp": "Stars and review count read from the Yelp listing.",
    "reviews.facebook": "Stars and review count read from the Facebook page.",
    "directories": "21 fixed directories searched by business name, city and site; each listing compared with the website's name, address and phone. Error rate = (Inaccurate + Not Found) / 21; Unchecked directories are not counted as errors.",
    "keyword_groups": "Services the prospect stressed on the call, phrased as customers search (no city), with SE Ranking US monthly volume. Volumes are national, not per city.",
    "rankings": "One city-targeted SE Ranking SERP per keyword and city, read for the local pack (positions 1 to 3) and organic page one (positions 1 to 10) by domain or business name.",
    "observations": "Written from the findings above, tied to what the prospect said on the call where possible, biggest business impact first.",
}

HOW_SHEET = "How We Got These Results"
HOW_HEADER = ("Section", "Item", "Result", "How it was checked", "What we found (evidence)", "Source")


def _evidence(data, key):
    """Return (note, source) for a dotted key, reading both evidence[key] and sources[key]."""
    ev = data.get("evidence", {}) or {}
    src = data.get("sources", {}) or {}
    short = key.split(".", 1)[-1]
    raw = ev.get(key, ev.get(short, ""))
    note, source = raw, ""
    if isinstance(raw, dict):
        note, source = raw.get("note", ""), raw.get("source", "")
    source = source or src.get(key, src.get(short, ""))
    return str(note or ""), str(source or "")


def _label_for(section_rows, key):
    return next((label for label, k in section_rows.items() if k == key), key)


def how_rows(data):
    rows = []

    def add(section, item, result, key, note="", source=""):
        n, s = _evidence(data, key)
        rows.append((section, item, "" if result is None else str(result), METHODS.get(key, ""), note or n, source or s))

    p = data.get("prospect", {})
    for k in ("company_name", "url", "audit_date"):
        add("Header", k.replace("_", " ").title(), p.get(k, ""), f"prospect.{k}")
    for label, key in SEO_ROWS.items():
        add("Getting Found: SEO & Content", label, data.get("seo", {}).get(key), f"seo.{key}")
    for label, key in UI_ROWS.items():
        add("Converting Visitors: UI & Performance", label.lstrip("- "), data.get("ui", {}).get(key), f"ui.{key}")
    for label, (section, key) in UNDER_HEADER.items():
        title = {"content": "Content", "ai_visibility": "Generative AI Visibility", "tracking": "Analytics & Ads"}[section]
        add(title, label, data.get(section, {}).get(key), f"{section}.{key}")
    add("Generative AI Visibility", "Other", data.get("ai_visibility", {}).get("other"), "ai_visibility.other")
    for label, key in REVIEW_COLUMNS.items():
        rev = data.get("reviews", {}).get(key, {})
        unknown = "unknown" in (str(rev.get("stars")).lower(), str(rev.get("count")).lower())
        add("GBP & Reviews", label, "Unknown" if unknown else f"{rev.get('stars', '')} stars, {rev.get('count', '')} reviews", f"reviews.{key}",
            note=str(rev.get("note", "")), source=str(rev.get("url", "")))

    d = data.get("directories", {})
    inaccurate, not_found = d.get("inaccurate") or 0, d.get("not_found") or 0
    add("Directory Listings", "Summary",
        f"{inaccurate} inaccurate, {not_found} not found, {round(100 * (inaccurate + not_found) / DIRECTORY_TOTAL)}% error rate",
        "directories")
    for item in d.get("detail", []):
        rows.append(("Directory Listings", item.get("directory", ""), item.get("status", ""),
                     "Listing searched and compared with the website's name, address and phone.",
                     item.get("note", ""), item.get("url", "")))

    for i, g in enumerate(data.get("keyword_groups", []), start=1):
        rows.append(("Keyword Groups", f"{i}. {g.get('keyword', '')}",
                     f"{g.get('searches_per_month', '')} searches/month, paired with {g.get('city', '')}",
                     METHODS["keyword_groups"], g.get("reason", ""), g.get("source", "SE Ranking keyword metrics, US")))

    for row in data.get("rankings", []):
        m, o = row.get("maps_position"), row.get("organic_position")
        rows.append(("Search Rankings", f"{row.get('keyword', '')} / {row.get('city', '')}",
                     f"Maps {m if m else '-'}, Organic {o if o else '-'}", METHODS["rankings"],
                     row.get("note", ""), row.get("source", "")))

    basis = data.get("observation_basis", []) or []
    for i, text in enumerate(data.get("observations", []), start=1):
        rows.append(("Key Observations", f"Observation {i}", text, METHODS["observations"],
                     basis[i - 1] if i - 1 < len(basis) else "", ""))
    return rows


def add_how_sheet(wb, data):
    ws = wb.create_sheet(HOW_SHEET)
    ws.append(HOW_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    last_section = None
    for row in how_rows(data):
        ws.append(row)
        r = ws.max_row
        for cell in ws[r]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0] != last_section:
            ws.cell(row=r, column=1).font = Font(bold=True)
            last_section = row[0]
        fill_ = colour(row[1], row[2])
        if fill_:
            ws.cell(row=r, column=3).fill = fill_
    for col, width in zip("ABCDEF", (28, 34, 30, 60, 60, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return ws


def save(grid, out_xlsx, data=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "WISE Digital Audit"
    bold_labels = {norm(k) for k in list(SEO_ROWS) + list(UI_ROWS) + list(UNDER_HEADER)}
    for r, row in enumerate(grid, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            if value != "" and value.lstrip("-").replace(".", "", 1).isdigit():
                cell.value = float(value) if "." in value else int(value)
            else:
                cell.value = value
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            left = norm(row[c - 3]) if c >= 3 else ""
            fill_ = colour(left, value)
            if fill_:
                cell.fill = fill_
            if norm(value) in bold_labels or norm(value).endswith(("Content", "Performance", "Result")):
                cell.font = Font(bold=True)
    for col, width in zip("ABCDEFG", (34, 16, 18, 4, 40, 16, 18)):
        ws.column_dimensions[col].width = width
    if data is not None:
        add_how_sheet(wb, data)
    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    with open(Path(out_xlsx).with_suffix(".csv"), "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(grid)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("audit_json")
    ap.add_argument("output_xlsx")
    ap.add_argument("--template", default=str(Path(__file__).resolve().parent.parent / "assets" / "audit-template.csv"))
    args = ap.parse_args()
    data = json.loads(Path(args.audit_json).read_text(encoding="utf-8"))
    grid = load_grid(args.template)
    problems = []
    fill(grid, data, problems)
    save(grid, args.output_xlsx, data)
    if problems:
        print("Saved with problems to fix:\n- " + "\n- ".join(problems), file=sys.stderr)
        sys.exit(1)
    print(f"Saved {args.output_xlsx} and {Path(args.output_xlsx).with_suffix('.csv')}")


if __name__ == "__main__":
    main()
